#!/usr/bin/env python3
"""Generate the investigator sign-off package (investigator_signoff/).

Populates sign-off sheets from the PROPOSED AI-assisted independent-review
records and the frozen mechanical outcomes. The investigator decision /
comment / initials / date fields are left EMPTY — this package requests, but
does not record, investigator approval.
"""
import csv
import json
from pathlib import Path

import xlsxwriter
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
ADJ = ROOT / "data/adjudication/main-study-20260925/verification_adjusted"
RUNS = ROOT / "results/run_outcomes.csv"
OUT = ROOT / "investigator_signoff"
OUT.mkdir(exist_ok=True)

SIGNOFF_FIELDS = [
    "investigator_decision(approve/reject/modify/pending)",
    "investigator_comment",
    "investigator_initials_or_name",
    "investigator_datetime_utc",
]


def read(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def write_csv(name, fieldnames, rows):
    with open(OUT / name, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)


frozen = {r["run_id"]: r for r in read(RUNS)}

a_rows = []
for r in read(ADJ / "PaperII_A_30slots_independent_review_PROPOSED.csv"):
    run_id = f"{r['paper_id'].split(':')[1]}-slot{r['slot']}"
    fz = frozen.get(run_id, {})
    frozen_label = "/".join(
        [fz.get(c, "") for c in ("L1_specification", "L2_implementation",
                                "L3_execution", "L4_numerical", "L5_conclusion")]
    ) + f"|{fz.get('failure_code', '')}"
    a_rows.append({
        "record_id": f"{r['paper_id']}/slot{r['slot']}",
        "frozen_classification": frozen_label,
        "proposed_adjusted_classification": (
            f"run_success={r['run_success']}; agree_with_mechanical={r['agree_with_mechanical']}"
        ),
        "evidence_ref": r["evidence_ref"],
        "reason_rationale": r["rationale"],
        **{k: "" for k in SIGNOFF_FIELDS},
    })
write_csv("A_30slots_signoff.csv",
          ["record_id", "frozen_classification", "proposed_adjusted_classification",
           "evidence_ref", "reason_rationale"] + SIGNOFF_FIELDS, a_rows)

b_rows = [{
    "record_id": r["paper_id"],
    "frozen_classification": r["recorded_state"],
    "proposed_adjusted_classification": (
        f"{r['corrected_state']} ({r['verification']})"
        if r["corrected_state"] else r["verification"]
    ),
    "evidence_ref": r["evidence_ref"],
    "reason_rationale": r["rationale"],
    **{k: "" for k in SIGNOFF_FIELDS},
} for r in read(ADJ / "PaperII_B_90papers_investigator_verification_PROPOSED.csv")]
write_csv("B_90papers_signoff.csv",
          ["record_id", "frozen_classification", "proposed_adjusted_classification",
           "evidence_ref", "reason_rationale"] + SIGNOFF_FIELDS, b_rows)

c_rows = [{
    "record_id": f"{r['paper_id']}|{r['layer']}",
    "frozen_classification": r["recorded_state"],
    "proposed_adjusted_classification": r["corrected_state"] or r["verification"],
    "evidence_ref": r["evidence_ref"],
    "reason_rationale": r["rationale"],
    **{k: "" for k in SIGNOFF_FIELDS},
} for r in read(ADJ / "PaperII_C_G1G5_pilot_candidate_verification_PROPOSED.csv")]
write_csv("C_G1G5_pilot_signoff.csv",
          ["record_id", "frozen_classification", "proposed_adjusted_classification",
           "evidence_ref", "reason_rationale"] + SIGNOFF_FIELDS, c_rows)

e_rows = [{
    "record_id": f"{r['paper_id']}|{r['dimension']}",
    "frozen_classification": r["recorded_completeness"],
    "proposed_adjusted_classification": (
        f"{r['corrected_completeness']} (agree={r['agree']})"
        if r["corrected_completeness"] else f"agree={r['agree']}"
    ),
    "evidence_ref": r["evidence_ref"],
    "reason_rationale": r["rationale"],
    **{k: "" for k in SIGNOFF_FIELDS},
} for r in read(ADJ / "PaperII_E_reveal_review_PROPOSED.csv")]
write_csv("E_reveal_signoff.csv",
          ["record_id", "frozen_classification", "proposed_adjusted_classification",
           "evidence_ref", "reason_rationale"] + SIGNOFF_FIELDS, e_rows)


def summary(name, title, n, rows_desc):
    (OUT / name).write_text(
        f"# {title}\n\n"
        f"Rows: {n}. Source: PROPOSED AI-assisted independent review under "
        f"`data/adjudication/main-study-20260925/verification_adjusted/`.\n"
        f"Status: provisional — NO investigator approval recorded. The\n"
        f"decision/comment/initials/datetime columns are intentionally empty\n"
        f"for the investigator to complete.\n\n{rows_desc}\n",
        encoding="utf-8",
    )


summary("A_SIGNOFF_SUMMARY.md", "Section A — 30 sealed slot records", len(a_rows),
        "Each row is one blind reconstruction slot. The frozen mechanical\n"
        "L1–L5/failure-code label is shown beside the proposed adjusted label.")
summary("B_SIGNOFF_SUMMARY.md", "Section B — 90 non-executed papers", len(b_rows),
        "Recorded barrier state vs proposed corrected state (incl. 81 verified\n"
        "barriers, 7 corrected, 2 unresolved).")
summary("C_SIGNOFF_SUMMARY.md", "Section C — G1–G5 pilot/candidate records", len(c_rows),
        "Second-pass confirmations and corrections of delegated G1–G5 states.")
summary("E_SIGNOFF_SUMMARY.md", "Section E — reveal-completeness records", len(e_rows),
        "Per-dimension recorded vs proposed corrected completeness labels.")

doc = Document()
doc.add_heading("Paper II — Investigator sign-off form (PRE-D)", 0)
doc.add_paragraph(
    "This form accompanies investigator_signoff/*.csv. Review each proposed "
    "record; enter approve / reject / modify / pending, a comment, your "
    "initials or name, and UTC date-time in the empty columns. Signing this "
    "form converts the PROPOSED AI-assisted independent review into "
    "investigator-adjudicated records. It does NOT constitute Human "
    "Validation D, which remains pending.")
doc.add_heading("Investigator", level=1)
for line in ("Name: ______________________", "Initials: __________",
             "Role: ______________________", "Date/time (UTC): ________________"):
    doc.add_paragraph(line)
doc.add_heading("Layer-level attestations", level=1)
for layer, n in (("A — 30 sealed slots", len(a_rows)),
                 ("B — 90 non-executed papers", len(b_rows)),
                 ("C — G1–G5 pilot/candidate", len(c_rows)),
                 ("E — reveal completeness", len(e_rows))):
    doc.add_paragraph(
        f"{layer}: {n} rows reviewed.   "
        "Decision for the layer as a whole:  [ ] approved  [ ] approved with "
        "per-row modifications  [ ] rejected", style="List Bullet")
doc.add_heading("Declaration", level=1)
doc.add_paragraph(
    "I confirm that I have reviewed the records listed above and that the "
    "decisions entered in the sign-off CSVs are mine. Signature: "
    "______________________   Date: ____________")
doc.save(OUT / "INVESTIGATOR_SIGNOFF_FORM.docx")

wb = xlsxwriter.Workbook(str(OUT / "INVESTIGATOR_SIGNOFF_CHECKLIST.xlsx"))
bold = wb.add_format({"bold": True})
req = wb.add_format({"bg_color": "#FEF3C7"})
for sheet, rows in (("A_30slots", a_rows), ("B_90papers", b_rows),
                    ("C_G1G5_pilot", c_rows), ("E_reveal", e_rows)):
    ws = wb.add_worksheet(sheet)
    headers = list(rows[0]) if rows else []
    for c, h in enumerate(headers):
        ws.write(0, c, h, bold)
    for i, row in enumerate(rows, start=1):
        for c, h in enumerate(headers):
            ws.write(i, c, row[h], req if h in SIGNOFF_FIELDS else None)
    ws.freeze_panes(1, 0)
ws = wb.add_worksheet("instructions")
ws.write(0, 0, "Fill ONLY the four investigator_* columns (highlighted). "
               "approve / reject / modify / pending. Leave untouched rows as "
               "pending. Do not edit frozen or proposed columns.", bold)
wb.close()

(OUT / "INVESTIGATOR_SIGNOFF_README.md").write_text(
    "# Investigator sign-off package (pre-D)\n\n"
    "Purpose: convert the PROPOSED AI-assisted independent-review records into\n"
    "investigator-adjudicated records. **No approval has occurred.** The\n"
    "investigator_* columns are empty by design.\n\n"
    "## Contents\n"
    "- `A_30slots_signoff.csv` (30), `B_90papers_signoff.csv` (90),\n"
    "  `C_G1G5_pilot_signoff.csv` (25), `E_reveal_signoff.csv` (130)\n"
    "- Per-layer summaries `*_SIGNOFF_SUMMARY.md`\n"
    "- `INVESTIGATOR_SIGNOFF_FORM.docx` — printable attestation\n"
    "- `INVESTIGATOR_SIGNOFF_CHECKLIST.xlsx` — editable workbook\n\n"
    "## Procedure\n"
    "1. Review each row against its `evidence_ref` (private evidence package\n"
    "   required for package-relative refs).\n"
    "2. Enter approve / reject / modify / pending plus comment, initials,\n"
    "   UTC datetime.\n"
    "3. Sign the DOCX form. Return the completed CSVs; the verification layer\n"
    "   is then updated to `investigator_signed` status.\n\n"
    "This is not Human Validation D (separate, still pending).\n",
    encoding="utf-8")
print(f"sign-off package: A={len(a_rows)} B={len(b_rows)} C={len(c_rows)} E={len(e_rows)}")
