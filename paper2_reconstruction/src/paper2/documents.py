import json
import zipfile
from pathlib import Path

import matplotlib
from docx import Document
from docx.document import Document as WordDocument
from docx.shared import Inches, Pt, RGBColor
from matplotlib import pyplot as plt
from matplotlib.patches import FancyBboxPatch
from pptx import Presentation
from pptx.dml.color import RGBColor as SlideColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as SlideInches
from pptx.util import Pt as SlidePt

from paper2.build import ROOT
from paper2.core import Row, read_csv, sha256

STATUS = "PREPARATION DRAFT — NOT SUBMISSION READY"
TITLE = "From scientific description to independent computational reconstruction"
SUBTITLE = "Prospective Paper II study linked to the EPJ commercial-software corpus"
CAPTION = (
    "Figure 1. Conceptual Wet/Dry framework. Paper II focuses on RECONSTRUCT and "
    "observes EXECUTE and REPRODUCE. ROBUST belongs to future Paper III and is excluded. "
    "The two pathways are an analogy, not a literal equivalence or an empirical result."
)
STAGES = ["ACCESS", "RECONSTRUCT", "EXECUTE", "REPRODUCE", "ROBUST"]
STAGE_COLORS = ["#E2E8F0", "#0F766E", "#CCFBF1", "#CCFBF1", "#FEF3C7"]
WET = [
    "Methods",
    "Materials",
    "Environmental /\nprocedural specification",
    "Reconstruct\nexperiment",
    "Execute",
    "Measurement /\nconclusion",
]
DRY = [
    "Methods",
    "Data / resources",
    "Computational\nspecification",
    "Independent\nimplementation",
    "Execute",
    "Numerical result /\nconclusion",
]
REFERENCES = (
    "acm_terms",
    "corebench",
    "goodman_crossref",
    "nasem_metadata",
    "paper2code",
    "paperbench_arxiv",
    "parent_crossref",
    "replicationbench",
    "scienceagentbench_arxiv",
    "wet_crossref",
)


def style_document(document: WordDocument) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    document.core_properties.author = "Paper II preparation; authorship not finalized"
    document.core_properties.subject = STATUS
    footer = section.footer.paragraphs[0]
    footer.text = STATUS
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string("9A3412")


def heading(document: WordDocument, title: str, level: int = 1) -> None:
    document.add_heading(title, level=level)


def paragraph(document: WordDocument, text: str) -> None:
    document.add_paragraph(text)


def caption(document: WordDocument, text: str) -> None:
    line = document.add_paragraph(text)
    line.paragraph_format.space_before = Pt(14)
    line.paragraph_format.space_after = Pt(8)
    line.paragraph_format.keep_with_next = text.startswith("Table")
    line.runs[0].italic = True


def table(document: WordDocument, rows: list[Row], columns: list[tuple[str, str]]) -> None:
    result = document.add_table(rows=1, cols=len(columns))
    result.style = "Table Grid"
    for cell, (_, label) in zip(result.rows[0].cells, columns, strict=False):
        cell.text = label
    for row in rows:
        for cell, (key, _) in zip(result.add_row().cells, columns, strict=False):
            cell.text = row[key].replace("_", " ")
    for row_index, table_row in enumerate(result.rows):
        for cell in table_row.cells:
            for line in cell.paragraphs:
                line.paragraph_format.keep_with_next = row_index < len(result.rows) - 1
                for run in line.runs:
                    run.font.size = Pt(9)


def markdown(document: WordDocument, text: str) -> None:
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("|"):
            cells = []
            while index < len(lines) and lines[index].startswith("|"):
                cells.append([cell.strip() for cell in lines[index].strip("|").split("|")])
                index += 1
            columns = [(str(column), name) for column, name in enumerate(cells[0])]
            rows = [{str(column): cell for column, cell in enumerate(row)} for row in cells[2:]]
            table(document, rows, columns)
            continue
        if line.startswith("# "):
            heading(document, line[2:], 1)
        elif line.startswith("## "):
            heading(document, line[3:], 2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            paragraph(document, line.strip("*"))
        index += 1


def framework(output: Path) -> None:
    matplotlib.rcParams["svg.fonttype"] = "none"
    matplotlib.rcParams["pdf.fonttype"] = 42
    figure, axis = plt.subplots(figsize=(13.333, 7.5))
    axis.set_xlim(0, 13.333)
    axis.set_ylim(0, 7.5)
    axis.axis("off")
    axis.text(
        0.2,
        7.1,
        "From description to a working scientific system",
        fontsize=21,
        weight="bold",
        color="#0F172A",
    )
    axis.text(0.2, 6.65, "Conceptual framework • Paper II preparation", fontsize=13)
    for index, (name, color) in enumerate(zip(STAGES, STAGE_COLORS, strict=False)):
        x = 0.2 + index * 2.6
        axis.add_patch(
            FancyBboxPatch(
                (x, 5.25),
                2.35,
                0.9,
                boxstyle="round,pad=0.03",
                facecolor=color,
                edgecolor="#64748B",
            )
        )
        axis.text(
            x + 1.175,
            5.7,
            name,
            ha="center",
            va="center",
            weight="bold",
            fontsize=14,
            color="white" if index == 1 else "#0F172A",
        )
        if index < len(STAGES) - 1:
            axis.text(x + 2.47, 5.7, "→", ha="center", va="center", fontsize=19)
    axis.text(4.0, 4.88, "Paper II focus and downstream observations", fontsize=12, color="#0F766E")
    axis.text(10.6, 4.88, "Paper III: excluded", fontsize=12, color="#92400E")
    for name, items, y, fill in (("WET", WET, 3.3, "#EFF6FF"), ("DRY", DRY, 1.7, "#F0FDFA")):
        axis.text(0.2, y + 1.08, name, fontsize=14, weight="bold", color="#0F172A")
        for index, label in enumerate(items):
            x = 0.2 + index * 2.17
            axis.add_patch(
                FancyBboxPatch(
                    (x, y),
                    1.95,
                    0.88,
                    boxstyle="round,pad=0.02",
                    facecolor=fill,
                    edgecolor="#CBD5E1",
                )
            )
            axis.text(x + 0.975, y + 0.44, label, fontsize=10.5, ha="center", va="center")
            if index < len(items) - 1:
                axis.text(x + 2.06, y + 0.44, "→", ha="center", va="center", fontsize=16)
    axis.text(
        0.2,
        0.8,
        "Independent implementation ≠ access to the original code",
        fontsize=13,
        color="#0F172A",
    )
    axis.text(
        0.2,
        0.35,
        "Conceptual analogy only. No robustness or multiverse analysis.",
        fontsize=12,
        color="#475569",
    )
    figure.tight_layout()
    for suffix in ("svg", "pdf", "png"):
        figure.savefig(output / f"figure1_framework.{suffix}", dpi=300, facecolor="white")
    plt.close(figure)
    svg = output / "figure1_framework.svg"
    svg.write_text("\n".join(line.rstrip() for line in svg.read_text().splitlines()) + "\n")


def slides(output: Path, characteristics: list[Row]) -> None:
    presentation = Presentation()
    presentation.slide_width = SlideInches(13.333)
    presentation.slide_height = SlideInches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    def textbox(x: float, y: float, width: float, height: float, text: str, size: int) -> None:
        shape = slide.shapes.add_textbox(
            SlideInches(x),
            SlideInches(y),
            SlideInches(width),
            SlideInches(height),
        )
        frame = shape.text_frame
        frame.margin_left = 0
        frame.margin_right = 0
        frame.word_wrap = True
        frame.text = text
        for line in frame.paragraphs:
            line.font.size = SlidePt(size)
            line.font.color.rgb = SlideColor.from_string("0F172A")

    textbox(0.3, 0.2, 12.5, 0.5, "Figure 1. From description to a working scientific system", 24)
    for index, (name, color) in enumerate(zip(STAGES, STAGE_COLORS, strict=False)):
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            SlideInches(0.3 + index * 2.6),
            SlideInches(1.05),
            SlideInches(2.25),
            SlideInches(0.72),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = SlideColor.from_string(color[1:])
        shape.line.color.rgb = SlideColor.from_string("64748B")
        shape.text = name
        line = shape.text_frame.paragraphs[0]
        line.alignment = PP_ALIGN.CENTER
        line.font.size = SlidePt(17)
        line.font.bold = True
        line.font.color.rgb = SlideColor.from_string("FFFFFF" if index == 1 else "0F172A")
        if index < 4:
            textbox(2.58 + index * 2.6, 1.13, 0.3, 0.5, "→", 20)
    textbox(3.0, 1.88, 7.5, 0.45, "Paper II focus and downstream observations", 15)
    textbox(10.7, 1.88, 2.3, 0.5, "Paper III: excluded", 14)
    for name, items, y in (("WET", WET, 2.85), ("DRY", DRY, 4.5)):
        textbox(0.3, y - 0.5, 2.0, 0.4, name, 18)
        for index, label in enumerate(items):
            shape = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                SlideInches(0.3 + index * 2.16),
                SlideInches(y),
                SlideInches(1.91),
                SlideInches(0.95),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = SlideColor.from_string("F0FDFA")
            shape.line.color.rgb = SlideColor.from_string("CBD5E1")
            shape.text_frame.word_wrap = True
            shape.text = label
            for line in shape.text_frame.paragraphs:
                line.font.size = SlidePt(13)
                line.font.color.rgb = SlideColor.from_string("0F172A")
                line.alignment = PP_ALIGN.CENTER
            if index < 5:
                textbox(2.22 + index * 2.16, y + 0.28, 0.22, 0.4, "→", 16)
    textbox(0.3, 6.12, 12.5, 0.9, CAPTION, 13)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox(0.3, 0.2, 12.5, 0.6, "Table 1. Deposited source rows and PMID identities", 24)
    grid = slide.shapes.add_table(
        len(characteristics) + 1,
        3,
        SlideInches(0.5),
        SlideInches(1.2),
        SlideInches(12.0),
        SlideInches(4.8),
    ).table
    for col, label in enumerate(("Recorded EPJ field", "Source rows", "Unique PMIDs in field")):
        grid.cell(0, col).text = label
    for row_index, row in enumerate(characteristics, 1):
        for col, key in enumerate(("field", "source_rows", "unique_pmids_within_field")):
            grid.cell(row_index, col).text = row[key].replace("_", " ")
    for table_row in grid.rows:
        for cell in table_row.cells:
            for line in cell.text_frame.paragraphs:
                line.font.size = SlidePt(16)
    textbox(
        0.5,
        6.35,
        12.0,
        0.7,
        "Observed source audit only. Field memberships overlap; do not sum within-field "
        "unique counts as a unique-paper total.",
        15,
    )
    presentation.save(str(output / "Figures_editable_EN.pptx"))


def reference_text(row: Row) -> str:
    authors = row["authors"]
    if row["reference_id"] == "parent_crossref":
        authors = "Onishi T; Ikenoue T [name order follows supplied author specification]"
    elif row["reference_id"] == "wet_crossref":
        authors = "Onishi T [name order follows supplied author specification]"
    elif row["reference_id"] == "nasem_metadata":
        authors = "National Academies of Sciences, Engineering, and Medicine"
    return f"{authors} ({row['year']}). {row['title']}. {row['publication_status']}. {row['URL']}"


def build_documents() -> None:
    output = ROOT / "manuscript"
    output.mkdir(parents=True, exist_ok=True)
    values = {
        r["claim_id"]: int(r["value"]) for r in read_csv(ROOT / "results/manuscript_values.csv")
    }
    characteristics = read_csv(ROOT / "results/corpus_characteristics.csv")
    precision = read_csv(ROOT / "results/precision_planning.csv")
    references = read_csv(ROOT / "data/verified_references.csv")
    framework(output)
    slides(output, characteristics)
    document = Document()
    style_document(document)
    heading(document, TITLE, 0)
    paragraph(document, SUBTITLE)
    paragraph(document, STATUS)
    paragraph(document, "Authors, affiliations and corresponding author: NOT FINALIZED.")
    paragraph(document, "Working journal: EPJ Research Infrastructures.")
    heading(document, "Preparation abstract")
    paragraph(
        document,
        "Access to research artifacts does not establish whether the published scientific "
        "description supports an independent implementation. This prospective study will "
        "estimate publication-grounded computational reconstructability under a specified "
        "agent, information boundary and resource budget, without consulting the original "
        "implementation or purchasing study-specific commercial software. The source audit "
        f"identified {values['source_records']:,} deposited records representing "
        f"{values['unique_pmids']:,} distinct PMIDs, requiring an explicit paper-level frame "
        "decision. Existing software and availability variables are retained as historical "
        "text detections rather than validated access assessments. The proposed design "
        "separates eligibility, input and resource access, specification, implementation, "
        "execution, numerical agreement and target-linked conclusion preservation. Three "
        "fixed independent run slots per sampled paper will support majority, strict and "
        "permissive outcome definitions while retaining missing or contaminated slots as "
        "unresolved. Targets, tolerances and resource ceilings must be frozen prospectively. "
        "This document contains only preparation findings and proposed methods. No pilot, "
        "reconstruction experiment, human validation or descriptive original-code reveal "
        "has been completed. No reconstructability rate or empirical conclusion is available. "
        "A qualified information firewall, verified study inputs and real validation remain "
        "necessary before this work can support a completed-study submission.",
    )
    paragraph(
        document,
        "Keywords: computational reconstruction; reproducibility; scientific specification; "
        "research infrastructure; software accessibility; research agents",
    )
    heading(document, "1 Introduction")
    paragraph(
        document,
        "Computational reproducibility terminology varies across communities (Goodman et al., "
        "2016; National Academies of Sciences, Engineering, and Medicine, 2019). ACM's "
        "post-2020 terminology distinguishes execution of authors' artifacts from results "
        "obtained using independently developed artifacts (ACM Publications Board, 2020). "
        "We therefore use publication-grounded independent computational reconstruction "
        "as an operational description rather than proposing a universal definition.",
    )
    paragraph(
        document,
        "The EPJ antecedent examines commercial software and version accessibility "
        "(Onishi and Ikenoue, 2026). Those conditions motivate, but do not answer, whether "
        "the scientific description contains enough operational detail to implement the "
        "reported method independently. The Wet antecedent concerns environmental "
        "confounding in clonal systems (Onishi, 2026). Its extension here is a conceptual "
        "analogy: nominal materials or data may be available while operational specification "
        "remains incomplete. The Wet full text was not accessible in the preparation audit; "
        "no detailed empirical claim is imported from it.",
    )
    paragraph(
        document,
        "Independent paper-to-code reconstruction already has substantial precedent. "
        "PaperBench restricts original implementations but supplies author clarifications "
        "(Starace et al., 2025); ReplicationBench uses curated astrophysics tasks and supplied "
        "execution information (Ye et al., 2025). Paper2Code studies repository generation "
        "(Seo et al., 2026), while ScienceAgentBench evaluates curated scientific programming "
        "tasks (Chen et al., 2025). CORE-Bench instead supplies original code and data "
        "(Siegel et al., 2024; revised 2026). Their endpoints and information packages differ "
        "and must not be pooled as a publication-level reconstructability rate.",
    )
    paragraph(
        document,
        "The proposed contribution is a finite-corpus estimation study with an explicit "
        "publication-only information boundary and zero study-specific software expenditure. "
        "It makes no first-study or terminological novelty claim. Figure 1 shows the "
        "proposed framework and the exclusion of robustness experiments.",
    )
    document.add_picture(str(output / "figure1_framework.png"), width=Inches(6.65))
    document.paragraphs[-1].paragraph_format.keep_with_next = True
    caption(document, CAPTION)
    heading(document, "2 Prospective methods — not yet performed")
    for title, text in (
        (
            "2.1 Source frame and eligibility",
            "Preserve deposited source records and derive paper identity separately. G1 assesses "
            "whether a central claim depends on a computational procedure with a potentially "
            "objective output. G2 identifies the principal target; G3/G4 assess lawful input and "
            "zero-purchase resource access; G5 describes specification sufficiency. Neither access "
            "nor specification failures exclude otherwise eligible papers from the all-eligible "
            "policy estimand. Historical code/data flags cannot replace these assessments.",
        ),
        (
            "2.2 Sampling and target selection",
            "After source resolution, a diverse pilot will finalize operational rules. The "
            "main sample is approximately 100 papers, stratified using a deterministic disjoint "
            "assignment of the seven overlapping EPJ fields. Save within-field sampling "
            "probabilities and weights. Before each run, select one central target by the frozen "
            "hierarchy and encode its associated claim. No outcome-informed target replacement "
            "is permitted. Planning precision in the supplement is not empirical evidence.",
        ),
        (
            "2.3 Independent reconstruction intervention",
            "Use three clean, fixed slots per paper with the same target, package and permissions "
            "and equivalent model and resource settings. Permitted materials are the publication, "
            "methodological supplement, vetted inputs and general documentation. Original "
            "implementations and other runs are prohibited. The broker must restrict model-side "
            "tools and execution networking. Separate VMs alone are insufficient. Prior model "
            "training exposure cannot be certified absent. No qualified harness currently exists.",
        ),
        (
            "2.4 Outcomes and comparison",
            "L1 specification, L2 implementation, L3 execution, L4 numerical reproduction and L5 "
            "conclusion preservation remain separate. A run succeeds only when an independent "
            "implementation executes and meets the frozen target-specific numerical rule. Majority "
            "paper success requires two clean successes among three fixed slots. Missing or "
            "contaminated slots remain unresolved, with bounds rather than reduced denominators. "
            "Exact counts, displayed-precision scalar comparisons and target-specific stochastic "
            "rules avoid arbitrary universal tolerances. Matching numbers do not by themselves "
            "prove implementation fidelity.",
        ),
        (
            "2.5 Statistical analysis and failure attribution",
            "Use design-weighted finite-corpus estimates and separately reported uncertainty and "
            "unknown-status bounds. Fixed-frame census counts do not have paper-sampling error. "
            "Report complete-triple consistency, hierarchical outcomes, failures and burden. "
            "Do not fit unjustified multivariable models or interpret code-sharing associations "
            "causally. Distinguish observed reconstruction failure from demonstrated publication "
            "inadequacy. The proposed weighted interval is not yet implemented; details and "
            "prespecified sensitivities are in the draft SAP.",
        ),
        (
            "2.6 Freeze, reveal and human validation",
            "Seal targets before runs and all outcomes before any original-code reveal. Reveal is "
            "descriptive only; no perturbation or revised score is allowed. Prepare approximately "
            "15 outcome-balanced human cases after agent freeze, with validators masked to agent "
            "and reveal results. Actual human work, institutional determination and participation "
            "records are indispensable. No human outcome is simulated. No robustness or multiverse "
            "analysis is part of Paper II.",
        ),
    ):
        heading(document, title, 2)
        paragraph(document, text)
    heading(document, "3 Observed preparation findings")
    paragraph(
        document,
        f"The two pinned source CSVs contain {values['source_records']:,} rows and "
        f"{values['unique_pmids']:,} unique PMIDs, with "
        f"{values['excess_records']:,} excess records over a unique-paper count. "
        f"There are {values['duplicated_pmids']:,} repeated identifiers: "
        f"{values['double_occurrence_pmids']:,} occur twice and "
        f"{values['triple_occurrence_pmids']:,} occurs three times. "
        f"The deposit has {values['unique_pmid_field_pairs']:,} distinct PMID-field "
        f"memberships, including {values['within_field_excess_records']:,} within-field "
        f"excess records and {values['multi_field_pmids']:,} PMIDs spanning fields. "
        "Non-field extracted values agree across repeated records. Table 1 presents "
        "source-row and within-field identity counts; it is not an eligibility or "
        "reconstruction-results table.",
    )
    caption(document, "Table 1. Observed deposited corpus structure.")
    table(
        document,
        characteristics,
        [
            ("field", "Recorded EPJ field"),
            ("source_rows", "Source rows"),
            ("unique_pmids_within_field", "Unique PMIDs within field"),
        ],
    )
    paragraph(
        document,
        "The original sampling code caps annual candidate retrieval and its nominal "
        "random-offset expression is always zero. Historical candidate/API snapshots and "
        "validation annotations were not recovered. Consequently original population-wide "
        "selection probabilities and extraction accuracy are not established by this "
        "audit. Current availability, eligibility and reconstruction outcomes remain "
        "unassessed. No reconstruction rate can be calculated.",
    )
    heading(document, "4 Discussion boundary and submission hold")
    paragraph(
        document,
        "The source audit supports identity and denominator statements only. It cannot "
        "establish that publications are reconstructable or unreconstructable. A completed "
        "study must separate infrastructure barriers, scientific specification and agent "
        "limitations; agent failure is not human impossibility. Original-code availability "
        "may correlate with reporting practice but cannot be interpreted causally here. "
        "Methods-linter development and robustness analysis remain future work.",
    )
    paragraph(
        document,
        "Submission is on hold for source-frame resolution, qualification of an enforced "
        "information firewall, pilot and prospective freeze, actual classification and "
        "reconstruction evidence, and human-validation status. This document must not be "
        "submitted as a completed empirical study or presented as a frozen protocol.",
    )
    heading(document, "Statements and declarations — pending author completion")
    paragraph(
        document,
        "Funding, competing interests, author contributions, affiliations and correspondence: "
        "not provided. Human-validation ethics/consent determination: not completed. "
        "Data/code: the public EPJ-linked repository contains the source and preparation "
        "pipeline; no Paper II reconstruction dataset exists. Third-party full text is "
        "retained internally under its applicable rights and is not redistributed.",
    )
    paragraph(
        document,
        "AI assistance: Devin performed source, literature and design audits and generated "
        "preparation code and draft materials. These preparation sessions are not blind "
        "reconstruction replicates. Human investigators must validate the design, factual "
        "claims, interpretation and final manuscript, and retain scientific responsibility.",
    )
    heading(document, "References used in this preparation draft")
    paragraph(
        document,
        "Antecedent author names follow the supplied author specification. Publisher/"
        "Crossref metadata reverse the first author's name components; verify the final "
        "bibliographic form before submission. Preprints and version-specific claims "
        "are identified in the reference ledger.",
    )
    chosen = [row for row in references if row["reference_id"] in REFERENCES]
    for row in sorted(chosen, key=reference_text):
        paragraph(document, reference_text(row))
    document.save(str(output / "Preparation_draft_EN.docx"))

    supplement = Document()
    style_document(supplement)
    heading(supplement, "Paper II — draft protocols and preparation supplement", 0)
    paragraph(supplement, STATUS)
    paragraph(
        supplement,
        "Table S1 gives analytic precision planning at assumed proportions. These values "
        "are not observed outcomes and do not replace design-based intervals.",
    )
    caption(supplement, "Table S1. Wilson planning intervals (unweighted binomial reference).")
    table(
        supplement,
        precision,
        [
            ("planned_n", "Planned n"),
            ("assumed_p", "Assumed p"),
            ("wilson_lower", "95% lower"),
            ("wilson_upper", "95% upper"),
            ("half_width", "Half-width"),
        ],
    )
    for path in sorted((ROOT / "protocols").glob("*.md")):
        supplement.add_page_break()
        markdown(supplement, path.read_text())
    supplement.save(str(output / "Supplement_protocols_DRAFT_EN.docx"))

    tables = Document()
    style_document(tables)
    heading(tables, "Editable preparation tables", 0)
    paragraph(tables, "Table 1 is the observed source audit; Table S1 is analytic planning only.")
    caption(tables, "Table 1. Observed deposited corpus structure.")
    table(
        tables,
        characteristics,
        [
            ("field", "Recorded EPJ field"),
            ("source_rows", "Source rows"),
            ("unique_pmids_within_field", "Unique PMIDs within field"),
        ],
    )
    caption(tables, "Table S1. Analytic Wilson precision planning.")
    table(
        tables,
        precision,
        [
            ("planned_n", "Planned n"),
            ("assumed_p", "Assumed p"),
            ("wilson_lower", "95% lower"),
            ("wilson_upper", "95% upper"),
            ("half_width", "Half-width"),
        ],
    )
    tables.save(str(output / "Editable_tables_EN.docx"))
    audit = (
        "# Repository audit — observed source only\n\n"
        f"Source records: {values['source_records']:,}; unique PMIDs: "
        f"{values['unique_pmids']:,}; excess records: {values['excess_records']:,}.\n\n"
        "The VOR links this deposited corpus. No corrected corpus or historical validation "
        "annotations were recovered in the audited public sources. Sampling candidates and "
        "original API snapshots are unavailable; historical flags are text-detection proxies. "
        "No replacement sample, blind reconstruction or human validation has occurred.\n\n"
        "See results/manuscript_values.csv for machine-readable numerical provenance, "
        "data/acquisition_ledger.csv for exact source hashes, review/EVIDENCE.md for retained "
        "audit evidence, and results/readiness.json for incomplete study gates.\n"
    )
    (ROOT / "review/REPOSITORY_AUDIT.md").write_text(audit)
    files = [
        *output.iterdir(),
        *sorted((ROOT / "protocols").glob("*.md")),
        *sorted((ROOT / "results").glob("*.csv")),
        *sorted((ROOT / "results").glob("*.json")),
        *sorted((ROOT / "src/paper2").glob("*.py")),
        *sorted((ROOT / "data/derived").glob("*.csv")),
        *sorted((ROOT / "data").glob("*.csv")),
        *sorted((ROOT / "schemas").glob("*")),
        ROOT / "requirements.lock",
        ROOT / "pyproject.toml",
    ]
    manifest = {
        str(path.relative_to(ROOT)): {
            "bytes": path.stat().st_size,
            "sha256": sha256(path.read_bytes()),
        }
        for path in files
        if path.is_file() and path.name != "preparation_artifact_manifest.json"
    }
    (ROOT / "results/preparation_artifact_manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n",
    )
    dist = ROOT / "dist"
    dist.mkdir(exist_ok=True)
    with zipfile.ZipFile(
        dist / "PaperII_preparation_NOT_SUBMISSION_READY.zip", "w", compression=zipfile.ZIP_DEFLATED
    ) as archive:
        for path in sorted(ROOT.rglob("*")):
            relative = path.relative_to(ROOT)
            if not path.is_file() or any(
                part.startswith(".") or part == "__pycache__" for part in relative.parts
            ):
                continue
            if relative.parts[0] in {"dist", "workflow"} or relative.parts[:2] == ("data", "raw"):
                continue
            archive.write(path, relative)


if __name__ == "__main__":
    build_documents()
