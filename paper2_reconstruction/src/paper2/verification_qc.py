"""QC for the verification-adjusted layer: frozen preservation, arithmetic, consistency."""

import json
import re
from pathlib import Path

from docx import Document

from paper2.core import read_csv, sha256
from paper2.main_analysis import mapping
from paper2.verification_layer import ADJUDICATION, BLIND, RESULTS, REVEAL_LEDGER, ROOT

FROZEN_BLIND_SHA = "350a0e7dcedbb74ccad2472c642318c87cc7d272f1b6c88186a1419dc858ea5a"
FROZEN_ADJUDICATION_SHA = "6fd42b8964e055606cbd24ebddb44f42fa09147fa80bfa2d5bcaa161766c5b28"
D_TERMS = ("validator_id", "consent_record", "human_results")


def docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def run() -> list[tuple[str, bool, str]]:
    audit = json.loads((RESULTS / "verification_import_audit.json").read_text())
    frozen = json.loads((ROOT / "results/main_analysis.json").read_text())
    values = {r["claim_id"]: r for r in read_csv(ROOT / "results/manuscript_values.csv")}
    est = mapping(audit["estimands"])
    cond = mapping(est["conditional_success_among_attempted"])
    b = mapping(audit["B"])
    funnel = mapping(b["funnel"])
    detector = mapping(b["leakage_detector_audit"])
    manuscript = docx_text(ROOT / "manuscript/Preparation_draft_EN.docx")
    tables_doc = docx_text(ROOT / "manuscript/Editable_tables_EN.docx")
    supplement = docx_text(ROOT / "manuscript/Supplement_protocols_DRAFT_EN.docx")
    slots = read_csv(RESULTS / "adjusted_A_slots.csv")
    papers = read_csv(RESULTS / "adjusted_A_papers.csv")
    dist = read_csv(RESULTS / "adjusted_A_distribution.csv")
    adjusted_success = sum(p["adjusted_majority"] == "success" for p in papers)
    frozen_success = sum(p["frozen_majority"] == "success" for p in papers)
    checks = [
        (
            "blind outcome byte-identical to seal",
            sha256(BLIND.read_bytes()) == FROZEN_BLIND_SHA,
            "",
        ),
        (
            "primary adjudication unchanged",
            sha256(ADJUDICATION.read_bytes()) == FROZEN_ADJUDICATION_SHA,
            "",
        ),
        (
            "reveal ledger references sealed blind outcome",
            json.loads(REVEAL_LEDGER.read_text())["blind_outcome_sha256"] == FROZEN_BLIND_SHA,
            "",
        ),
        (
            "frozen main_analysis conditional successes unchanged (0)",
            frozen["conditional_rate_among_attempted"]["successes"] == 0,
            "",
        ),
        (
            "A rows 30 / B 90 / C 25 / E 130",
            audit["rows"] == {"A": 30, "B": 90, "C": 25, "E": 130},
            "",
        ),
        (
            "no unresolved evidence references",
            mapping(audit["evidence_references"])["unresolved"] == 0,
            "",
        ),
        ("no validation errors", audit["validation_errors"] == [], ""),
        (
            "no reruns / no original code executed",
            b["reruns_performed"] == 0 and not mapping(audit["E"])["original_code_executed"],
            "",
        ),
        (
            "no investigator sign-off populated",
            all(r["investigator_signoff"] == "" for r in slots),
            "",
        ),
        (
            "adjusted conditional numerator equals paper table majority count",
            mapping(cond["verification_adjusted"])["numerator"] == adjusted_success
            and mapping(cond["frozen_mechanical"])["numerator"] == frozen_success,
            f"adjusted {adjusted_success}, frozen {frozen_success}",
        ),
        (
            "distribution sums to 10 papers (frozen and adjusted)",
            sum(int(r["frozen_papers"]) for r in dist) == 10
            and sum(int(r["adjusted_papers"]) for r in dist) == 10,
            "",
        ),
        (
            "funnel columns each sum to 100",
            all(
                sum(int(str(v)) for v in mapping(funnel[k]).values()) == 100
                for k in ("frozen", "adjusted", "envelope")
            ),
            "",
        ),
        (
            "B verification counts sum to 90",
            sum(int(str(b[k])) for k in ("confirmed", "state_incorrect", "unknown")) == 90,
            "",
        ),
        (
            "leakage audit partitions flags",
            sum(int(str(detector[k])) for k in ("confirmed", "corrected", "unresolved"))
            == int(str(detector["classifications"])),
            "",
        ),
        (
            "manuscript_values frozen/adjusted numerators match audit",
            values["frozen_majority_successes_among_attempted"]["value"]
            == str(mapping(cond["frozen_mechanical"])["numerator"])
            and values["adjusted_majority_successes_among_attempted"]["value"]
            == str(mapping(cond["verification_adjusted"])["numerator"]),
            "",
        ),
        (
            "manuscript narrative cites frozen and adjusted conditional values",
            f"{mapping(cond['frozen_mechanical'])['numerator']}/10 (" in manuscript
            and f"{mapping(cond['verification_adjusted'])['numerator']}/10 (" in manuscript,
            "",
        ),
        ("manuscript cites Tables 6-9", all(f"Table {n}" in manuscript for n in (6, 7, 8, 9)), ""),
        (
            "Tables 6-9 present in editable tables and supplement",
            all(f"Table {n}." in tables_doc and f"Table {n}." in supplement for n in (6, 7, 8, 9)),
            "",
        ),
        (
            "manuscript labels review as AI-assisted and provisional, D pending",
            "AI-assisted" in manuscript
            and "provisional" in manuscript
            and "Section D" in manuscript
            and "pending" in manuscript,
            "",
        ),
        (
            "manuscript does not claim submission readiness",
            "NOT SUBMISSION READY" in manuscript
            and not re.search(
                r"\bsubmission[- ]ready\b(?!\W*NOT)", manuscript.replace("NOT SUBMISSION READY", "")
            ),
            "",
        ),
        ("D case count is zero", values["D_human_validation_cases_completed"]["value"] == "0", ""),
        (
            "no D result terms in adjusted outputs",
            not any(t in (RESULTS / "verification_import_audit.json").read_text() for t in D_TERMS),
            "",
        ),
    ]
    return checks


def write_report() -> None:
    checks = run()
    lines = [
        "# Verification-adjusted layer QC report",
        "",
        "Generated by `src/paper2/verification_qc.py` from current `results/` and `manuscript/`.",
        "",
        "| Check | Result | Note |",
        "|---|---|---|",
    ]
    for name, ok, note in checks:
        lines.append(f"| {name} | {'PASS' if ok else 'FAIL'} | {note} |")
    passed = sum(ok for _, ok, _ in checks)
    lines += ["", f"{passed}/{len(checks)} checks passed.", ""]
    (ROOT / "review/VERIFICATION_ADJUSTED_QC.md").write_text("\n".join(lines))
    if passed != len(checks):
        raise SystemExit("verification QC failed")


if __name__ == "__main__":
    write_report()
