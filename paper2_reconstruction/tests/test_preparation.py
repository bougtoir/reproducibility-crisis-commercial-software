import json
from pathlib import Path

import pytest
from docx import Document

from paper2.build import EMPTY_TABLES, PINS, ROOT, prepare_experiment_tables
from paper2.core import read_csv, sha256, write_csv


def test_existing_observations_stop_preparation_without_changes(tmp_path: Path) -> None:
    path = tmp_path / "run_outcomes.csv"
    write_csv(
        path, [{"paper_id": "synthetic-software-test-only"}], EMPTY_TABLES["run_outcomes.csv"]
    )
    original = path.read_bytes()
    with pytest.raises(ValueError, match="must not overwrite"):
        prepare_experiment_tables(tmp_path)
    assert path.read_bytes() == original
    assert list(tmp_path.iterdir()) == [path]


def test_empty_tables_with_wrong_schema_are_not_silently_accepted(tmp_path: Path) -> None:
    path = tmp_path / "run_outcomes.csv"
    write_csv(path, [], ["wrong_schema"])
    with pytest.raises(ValueError, match="Unexpected experimental table schema"):
        prepare_experiment_tables(tmp_path)


def test_real_source_hashes_and_record_identity_bridge() -> None:
    for source, digest in PINS.items():
        assert sha256((ROOT / "data/raw/epj" / source).read_bytes()) == digest
    raw = read_csv(ROOT / "data/raw/epj/output/extracted_data.csv")
    bridge = read_csv(ROOT / "data/derived/source_record_bridge.csv")
    papers = read_csv(ROOT / "data/derived/paper_registry.csv")
    assert len(raw) == len(bridge)
    assert len({row["source_record_id"] for row in bridge}) == len(raw)
    assert {row["pmid"] for row in raw} == {row["pmid"] for row in papers}
    assert [(row["pmid"], row["stratum"]) for row in raw] == [
        (row["pmid"], row["epj_field"]) for row in bridge
    ]


def test_no_fabricated_observations_or_submission_readiness() -> None:
    for filename in EMPTY_TABLES:
        assert read_csv(ROOT / "data/derived" / filename) == []
    status = json.loads((ROOT / "results/readiness.json").read_text())
    assert status["primary_success_rate"] == "not_assessable"
    assert status["study_status"] == "PREPARATION_ONLY_NOT_SUBMISSION_READY"
    queue = read_csv(ROOT / "data/derived/funnel_NOT_ASSESSED.csv")
    assert {row["assessment_status"] for row in queue} == {"NOT_STARTED"}
    assert {row["G1_reason"] for row in queue} == {"not_assessed"}


def test_document_values_come_from_canonical_results() -> None:
    document = Document(str(ROOT / "manuscript/Preparation_draft_EN.docx"))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    values = {
        row["claim_id"]: int(row["value"])
        for row in read_csv(ROOT / "results/manuscript_values.csv")
    }
    for key in ("source_records", "unique_pmids", "excess_records", "duplicated_pmids"):
        assert f"{values[key]:,}" in text
    assert "NOT SUBMISSION READY" in text
    assert "Figure 1" in text and "Table 1" in text
    assert len(document.inline_shapes) == 1
    assert len(document.tables) == 1


def test_artifact_manifest_matches_current_bytes_and_has_no_self_reference() -> None:
    path = ROOT / "results/preparation_artifact_manifest.json"
    manifest = json.loads(path.read_text())
    assert str(path.relative_to(ROOT)) not in manifest
    for relative, properties in manifest.items():
        artifact: Path = ROOT / relative
        assert artifact.stat().st_size == properties["bytes"]
        assert sha256(artifact.read_bytes()) == properties["sha256"]
